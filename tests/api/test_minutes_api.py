"""Activities Dashboard API (minutes registers): registers, FY filter, access, writes, archive, audit, parse and Excel."""
import io
import zipfile
from xml.sax.saxutils import escape

import pytest
from openpyxl import load_workbook

BRIGHT = 2          # BrightMed Pharma: seeded with 3 board meetings and 1 AGM in FY 2024-25
NEXA = 1            # the read-only client's company


@pytest.fixture(scope="module", autouse=True)
def leave_no_meetings_behind(admin, owner):
    """The test database is shared by the whole session: archive whatever these tests create so other tests'
    exact meeting counts are unaffected (nothing is ever deleted)."""
    first = owner.execute("SELECT coalesce(max(id), 0) AS n FROM core.meeting").fetchone()["n"]
    yield
    for r in owner.execute("SELECT id FROM core.meeting WHERE id > %s AND NOT is_archived", (first,)).fetchall():
        admin.ok(admin.post(f"/entities/meeting/{r['id']}/archive"))


def docx(lines):
    paras = "".join(f"<w:p><w:r><w:t>{escape(l)}</w:t></w:r></w:p>" for l in lines)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", f"<w:document><w:body>{paras}</w:body></w:document>")
    return buf.getvalue()


def board(admin, fy="2024-25", cid=BRIGHT):
    return admin.ok(admin.get(f"/minutes/board?company_id={cid}&fy={fy}"))


def test_registers_and_counts(admin):
    regs = admin.ok(admin.get("/minutes/registers"))
    assert [r["key"] for r in regs["registers"]] == ["board", "audit", "nrc", "src", "csr", "rmc", "agm", "egm"]
    assert regs["modes"] == ["Physical", "Video conferencing", "OAVM"]
    c = admin.ok(admin.get(f"/minutes/counts?company_id={BRIGHT}&fy=2024-25"))["counts"]
    assert c["board"] == 3 and c["agm"] == 1 and c["egm"] == 0


def test_list_shape_and_compliance(admin):
    data = board(admin)
    assert len(data["meetings"]) == 3
    m = data["meetings"][0]
    for key in ("meeting_id", "register", "meeting_date", "mode", "agenda", "attendance", "updated_at", "comments"):
        assert key in m
    chips = {c["key"]: c for c in data["compliance"]}
    assert chips["min"]["status"] == "attention" and chips["min"]["detail"] == "3 recorded"
    assert admin.ok(admin.get(f"/minutes/board?company_id={BRIGHT}&fy=2023-24"))["meetings"] == []


def test_bad_inputs(admin):
    assert admin.get(f"/minutes/board?company_id={BRIGHT}&fy=2025-27").status_code == 422
    assert admin.get(f"/minutes/nope?company_id={BRIGHT}&fy=2025-26").status_code == 404


def test_company_access(client, manager):
    assert client.get(f"/minutes/board?company_id={BRIGHT}&fy=2024-25").status_code == 404
    assert client.get(f"/minutes/counts?company_id={BRIGHT}&fy=2024-25").status_code == 404
    assert client.get(f"/minutes/roster?company_id={BRIGHT}&date=2025-01-01").status_code == 404
    assert client.get(f"/minutes/export.xlsx?company_id={BRIGHT}&fy=2024-25").status_code == 404
    assert client.get(f"/minutes/board?company_id={NEXA}&fy=2025-26").status_code == 200
    assert manager.get(f"/minutes/board?company_id={NEXA}&fy=2025-26").status_code == 404      # not theirs


def test_read_only_user_cannot_write_or_parse(client):
    body = {"meeting_date": "2025-05-01"}
    assert client.post(f"/minutes/board?company_id={NEXA}", body).status_code == 403
    assert client.put(f"/minutes/board/1?company_id={NEXA}", body).status_code == 403
    r = client.http.post("/api/minutes/parse", data={"company_id": str(NEXA)},
                         files={"file": ("m.docx", docx(["x"]))})
    assert r.status_code == 403


def test_fy_boundary(admin, owner):
    a = admin.ok(admin.post(f"/minutes/board?company_id={BRIGHT}", {"meeting_date": "2026-03-31"}), 201)["meeting"]
    b = admin.ok(admin.post(f"/minutes/board?company_id={BRIGHT}", {"meeting_date": "2026-04-01"}), 201)["meeting"]

    def ids(fy):
        return [m["meeting_id"] for m in board(admin, fy)["meetings"]]

    assert a["meeting_id"] in ids("2025-26") and b["meeting_id"] not in ids("2025-26")
    assert b["meeting_id"] in ids("2026-27") and a["meeting_id"] not in ids("2026-27")
    row = owner.execute("SELECT financial_year FROM core.meeting WHERE id = %s", (b["meeting_id"],)).fetchone()
    assert row["financial_year"] == "2026-27"                         # derived from the date, never trusted from a client


def test_create_edit_archive_audit(admin, owner):
    roster = admin.ok(admin.get(f"/minutes/roster?company_id={BRIGHT}&date=2025-06-10"))["roster"]
    assert len(roster) >= 2
    d1, d2 = roster[0]["directorship_id"], roster[1]["directorship_id"]
    body = {"meeting_date": "2025-06-10", "mode": "Physical", "meeting_time": "11:00 AM to 12:30 PM",
            "notice_date": "2025-06-01", "chairman": roster[0]["full_name"], "minutes_status": "Signed",
            "signing_date": "2025-06-20", "signing_place": "Mumbai", "comments": "ok",
            "attendance": [{"directorship_id": d1, "attended": True}, {"directorship_id": d1, "attended": True},
                           {"directorship_id": d2, "attended": False}],
            "agenda": [{"text": "Approve accounts", "reg30": "Yes", "forms": "AOC-4", "docs": "Auditor report",
                        "comments": "c"},
                       {"text": "   ", "reg30": "NA"}]}
    m = admin.ok(admin.post(f"/minutes/board?company_id={BRIGHT}", body), 201)["meeting"]
    mid = m["meeting_id"]
    assert m["meeting_time"] == "11:00 AM to 12:30 PM" and m["mode"] == "Physical" and m["minutes_status"] == "Signed"
    assert m["notice_date"] == "2025-06-01" and m["signing_place"] == "Mumbai"
    assert m["chairman"] == roster[0]["full_name"]
    assert len(m["agenda"]) == 1 and m["agenda"][0]["reg30"] == "Yes" and m["agenda"][0]["forms"] == "AOC-4"
    assert m["directors_present"] == roster[0]["full_name"]
    n = owner.execute("SELECT count(*) AS n FROM core.meeting_attendance WHERE meeting_id = %s AND directorship_id = %s",
                      (mid, d1)).fetchone()["n"]
    assert n == 1                                                      # ticking a director twice leaves one row
    assert mid in [x["meeting_id"] for x in board(admin, "2025-26")["meetings"]]

    # edit: update the agenda row by id, add one, tick the second director, change the status
    ag = m["agenda"][0]
    body2 = {**body, "updated_at": m["updated_at"], "minutes_status": "Entered in Minutes Book",
             "attendance": [{"directorship_id": d1, "attended": True}, {"directorship_id": d2, "attended": True}],
             "agenda": [{**ag, "text": "Approve accounts (revised)"}, {"text": "Second item", "reg30": "NA"}]}
    e = admin.ok(admin.put(f"/minutes/board/{mid}?company_id={BRIGHT}", body2))["meeting"]
    assert e["minutes_status"] == "Entered in Minutes Book" and len(e["agenda"]) == 2
    assert e["agenda"][0]["id"] == ag["id"] and e["agenda"][0]["text"] == "Approve accounts (revised)"
    assert e["attendance_present"] == 2
    # a stale save is refused
    stale = admin.put(f"/minutes/board/{mid}?company_id={BRIGHT}", {**body2, "updated_at": m["updated_at"]})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "stale"
    # dropping an agenda row archives it; it is never deleted
    cur = admin.ok(admin.put(f"/minutes/board/{mid}?company_id={BRIGHT}",
                             {**body2, "updated_at": e["updated_at"], "agenda": [e["agenda"][0]]}))["meeting"]
    assert len(cur["agenda"]) == 1
    arch = owner.execute("SELECT count(*) AS n FROM core.meeting_agenda_item WHERE meeting_id = %s AND is_archived",
                         (mid,)).fetchone()["n"]
    assert arch == 1

    # field errors
    bad = admin.put(f"/minutes/board/{mid}?company_id={BRIGHT}", {"meeting_date": "2025-06-10", "mode": "Telepathy",
                                                                  "notice_date": "31/12", "meeting_time": "late"})
    assert bad.status_code == 422 and set(bad.json()["error"]["fields"]) >= {"mode", "notice_date", "meeting_time"}

    # archive hides but keeps; the audit log names the actor
    admin.ok(admin.post(f"/entities/meeting/{mid}/archive"))
    assert mid not in [x["meeting_id"] for x in board(admin, "2025-26")["meetings"]]
    assert owner.execute("SELECT is_archived FROM core.meeting WHERE id = %s", (mid,)).fetchone()["is_archived"]
    actors = {r["actor"] for r in owner.execute(
        "SELECT actor FROM audit.log WHERE table_name = 'meeting' AND record_id = %s", (mid,)).fetchall()}
    assert actors == {admin.user["id"]}
    assert owner.execute("SELECT count(*) AS n FROM audit.log WHERE table_name = 'meeting_attendance' AND actor = %s",
                         (admin.user["id"],)).fetchone()["n"] > 0


def test_director_must_be_on_the_board_that_day(admin):
    r = admin.post(f"/minutes/board?company_id={BRIGHT}", {"meeting_date": "2025-06-11",
                                                          "attendance": [{"directorship_id": 999999, "attended": True}]})
    assert r.status_code == 422 and "attendance" in r.json()["error"]["fields"]


def test_wrong_register_or_company_for_update(admin, manager):
    m = admin.ok(admin.post(f"/minutes/board?company_id={BRIGHT}", {"meeting_date": "2025-07-01"}), 201)["meeting"]
    assert admin.put(f"/minutes/agm/{m['meeting_id']}?company_id={BRIGHT}", {"meeting_date": "2025-07-01"}).status_code == 404
    assert admin.put(f"/minutes/board/{m['meeting_id']}?company_id=3", {"meeting_date": "2025-07-01"}).status_code == 404
    assert manager.put(f"/minutes/board/{m['meeting_id']}?company_id=1", {"meeting_date": "2025-07-01"}).status_code == 404


def test_committee_registers(admin):
    # BrightMed has an Audit Committee but no NRC: Add is refused with a clear reason
    made = admin.ok(admin.post(f"/minutes/audit?company_id={BRIGHT}", {"meeting_date": "2025-08-01"}), 201)["meeting"]
    assert made["register"] == "audit"
    r = admin.post(f"/minutes/nrc?company_id={BRIGHT}", {"meeting_date": "2025-08-01"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "no_committee"
    assert admin.ok(admin.get(f"/minutes/nrc?company_id={BRIGHT}&fy=2025-26"))["committee_available"] is False
    assert admin.ok(admin.get(f"/minutes/audit?company_id={BRIGHT}&fy=2025-26"))["committee_available"] is True
    assert admin.ok(admin.get(f"/minutes/board?company_id={BRIGHT}&fy=2025-26"))["committee_available"] is True


def test_parse_endpoint(admin):
    lines = ["Held on 15.04.2025", "1. To approve the appointment of the statutory auditor of the company"]
    r = admin.http.post("/api/minutes/parse", data={"company_id": str(BRIGHT)}, files={"file": ("BM.docx", docx(lines))})
    data = r.json()
    assert r.status_code == 200 and data["source"] == "BM.docx" and data["meeting_date"] == "2025-04-15"
    assert data["agenda"][0]["forms"] == ["ADT-1"]
    bad = admin.http.post("/api/minutes/parse", data={"company_id": str(BRIGHT)}, files={"file": ("n.txt", b"hi")})
    assert bad.status_code == 400
    import pymupdf
    d = pymupdf.open()
    d.new_page()
    scan = admin.http.post("/api/minutes/parse", data={"company_id": str(BRIGHT)}, files={"file": ("s.pdf", d.tobytes())})
    assert scan.status_code == 422 and "scanned" in scan.json()["error"]["message"]


def test_excel_export(admin, client):
    r = admin.get(f"/minutes/export.xlsx?company_id={BRIGHT}&fy=2024-25")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert r.headers["content-disposition"].startswith('attachment; filename="Minutes_BrightMed_Pharma_Limited_2024-25')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["BM", "AGM"] and wb["BM"].max_row == 5
    one = load_workbook(io.BytesIO(admin.get(f"/minutes/export.xlsx?company_id={BRIGHT}&fy=2024-25&register=agm").content))
    assert one.sheetnames == ["AGM"]
    empty = load_workbook(io.BytesIO(admin.get(f"/minutes/export.xlsx?company_id={BRIGHT}&fy=2019-20").content))
    assert empty.sheetnames == ["Meetings"]
    assert client.get(f"/minutes/export.xlsx?company_id={NEXA}&fy=2025-26").status_code == 200
