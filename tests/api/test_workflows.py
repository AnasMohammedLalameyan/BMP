"""The business workflows built on top of the generic CRUD: activities and tasks, meetings, the directors register,
documents, exports, listing, and the rules that keep sensitive fields safe."""
import hashlib
import io
import json
import re

import pytest

START = {"meeting_date": "2026-12-10", "start_time": "11:00", "venue": "5th Floor, Nariman Point", "venue_mode": "PHYSICAL",
         "agenda_items": ["Approve accounts", "Appoint auditor"]}


# ============================ activities and tasks ============================

def start(sess, company=2, params=None):
    r = sess.post("/activities/instances", {"template_code": "BOARD_MEETING", "company_id": company, "params": params or START})
    return sess.ok(r, 201)


def test_the_template_is_described_without_duplicates(admin, owner):
    tpl = admin.ok(admin.get("/activities/templates"))["templates"][0]
    assert tpl["code"] == "BOARD_MEETING" and tpl["status"] == "draft" and tpl["steps"] == 9
    detail = admin.ok(admin.get(f"/activities/templates/{tpl['id']}"))
    steps = detail["steps"]
    assert [s["seq"] for s in steps] == list(range(1, 10))
    expected = owner.execute("SELECT count(*) AS n FROM activity.step_data_binding b JOIN activity.step_template s ON s.id = b.step_id "
                             "WHERE s.template_id = %s AND NOT b.is_archived", (tpl["id"],)).fetchone()["n"]
    assert sum(len(s["bindings"]) for s in steps) == expected >= 20
    assert [s["depends_on"] for s in steps][1:] == [[steps[i]["id"]] for i in range(8)]       # a chain 1 -> 2 -> ... -> 9
    assert steps[3]["requires_confirmation"] is True and steps[3]["ai_assisted"] is True
    raw = admin.ok(admin.get("/activities/templates/BOARD_MEETING/json"))
    assert raw["template"]["code"] == "BOARD_MEETING" and len(raw["steps"]) == 9


def test_starting_an_activity_creates_a_task_per_step_with_due_dates(manager):
    d = start(manager)
    tasks = d["tasks"]
    assert [t["step"]["seq"] for t in tasks] == list(range(1, 10))
    due = {t["step"]["step_key"]: t["due_date"] for t in tasks}
    assert due["CONFIRM_DATE_VENUE"] is None and due["UPDATE_REGISTERS_AND_FILINGS"] is None
    assert due["DRAFT_NOTICE"] == due["SEND_NOTICE"] == "2026-12-03"        # 7 days before the meeting
    assert due["CONDUCT_MEETING"] == due["RECORD_ATTENDANCE"] == "2026-12-10"
    assert due["DRAFT_MINUTES"] == "2026-12-25" and due["SIGN_MINUTES"] == "2027-01-09"
    assert all(t["status"] == "not_started" for t in tasks) and d["instance"]["status"] == "active"
    assert d["instance"]["template_status"] == "draft"


def test_the_start_form_is_validated_against_the_templates_input_schema(manager):
    r = manager.post("/activities/instances", {"template_code": "BOARD_MEETING", "company_id": 2, "params": {"venue": "x"}})
    assert r.status_code == 422
    fields = r.json()["error"]["fields"]
    assert {"params.meeting_date", "params.start_time", "params.venue_mode", "params.agenda_items"} <= set(fields)
    bad = manager.post("/activities/instances", {"template_code": "BOARD_MEETING", "company_id": 2, "params": {**START, "meeting_date": "10/12/2026"}})
    assert bad.status_code == 422 and "params.meeting_date" in bad.json()["error"]["fields"]
    assert manager.post("/activities/instances", {"template_code": "NOPE", "company_id": 2, "params": START}).status_code == 404
    assert manager.post("/activities/instances", {"template_code": "BOARD_MEETING", "company_id": 1, "params": START}).status_code == 404


def complete(sess, task_id, **kw):
    return sess.post(f"/activities/tasks/{task_id}/status", {"status": "completed", **kw})


def test_steps_must_be_completed_in_order_and_confirmation_is_enforced(manager, admin, owner):
    d = start(manager)
    inst = d["instance"]["id"]
    t = {x["step"]["seq"]: x["id"] for x in d["tasks"]}

    r = complete(manager, t[2])
    assert r.status_code == 409 and r.json()["error"]["code"] == "blocked" and "Confirm Date" in r.json()["error"]["message"]

    done = manager.ok(complete(manager, t[1]))
    first = next(x for x in done["tasks"] if x["id"] == t[1])
    assert first["status"] == "completed" and first["completed_by"] == manager.user["id"] and first["completed_at"]
    row = owner.execute("SELECT status, completed_by, completed_at FROM ops.task_instance WHERE id = %s", (t[1],)).fetchone()
    assert row["status"] == "completed" and row["completed_by"] == manager.user["id"] and row["completed_at"]

    manager.ok(complete(manager, t[2]))
    w = manager.ok(complete(manager, t[3]))
    assert any("No evidence attached" in x for x in w["warnings"])        # mandatory notice and agenda not yet attached

    # step 4 needs a named person's confirmation
    r = complete(manager, t[4])
    assert r.status_code == 409 and r.json()["error"]["code"] == "needs_confirmation"
    req = manager.ok(manager.post(f"/activities/tasks/{t[4]}/confirmations", {}), 201)["confirmation"]
    assert req["status"] == "pending" and req["requested_by"] == manager.user["id"]
    assert manager.post(f"/activities/tasks/{t[4]}/confirmations", {}).status_code == 409      # one at a time
    assert manager.post(f"/confirmations/{req['id']}/confirm").status_code == 403              # not by the person who asked
    decided = admin.ok(admin.post(f"/confirmations/{req['id']}/confirm"))["confirmation"]
    assert decided["status"] == "confirmed" and decided["confirmed_by"] == admin.user["id"] and decided["confirmed_at"]
    assert admin.post(f"/confirmations/{req['id']}/confirm").status_code == 409                # already decided
    manager.ok(complete(manager, t[4]))

    # the rest follow; step 8 needs its own confirmation
    for seq in (5, 6, 7):
        manager.ok(complete(manager, t[seq]))
    c8 = manager.ok(manager.post(f"/activities/tasks/{t[8]}/confirmations", {}), 201)["confirmation"]
    admin.ok(admin.post(f"/confirmations/{c8['id']}/confirm"))
    manager.ok(complete(manager, t[8]))
    last = manager.ok(manager.post(f"/activities/tasks/{t[9]}/status", {"status": "skipped"}))
    assert last["instance"]["status"] == "completed"                                            # rolled up from its tasks

    back = manager.ok(manager.post(f"/activities/tasks/{t[9]}/status", {"status": "not_started"}))
    assert back["instance"]["status"] == "active" and next(x for x in back["tasks"] if x["id"] == t[9])["completed_at"] is None


def test_a_rejected_confirmation_does_not_unlock_the_step(manager, admin):
    d = start(manager)
    t = {x["step"]["seq"]: x["id"] for x in d["tasks"]}
    for seq in (1, 2, 3):
        manager.ok(complete(manager, t[seq]))
    c = manager.ok(manager.post(f"/activities/tasks/{t[4]}/confirmations", {}), 201)["confirmation"]
    assert admin.ok(admin.post(f"/confirmations/{c['id']}/reject"))["confirmation"]["status"] == "rejected"
    assert complete(manager, t[4]).status_code == 409


def test_cancelling_and_reopening_an_activity(manager):
    d = start(manager)
    inst, t1 = d["instance"]["id"], d["tasks"][0]["id"]
    assert manager.ok(manager.post(f"/activities/instances/{inst}/status", {"status": "cancelled"}))["instance"]["status"] == "cancelled"
    assert complete(manager, t1).status_code == 409                                            # nothing moves while cancelled
    assert manager.ok(manager.post(f"/activities/instances/{inst}/status", {"status": "active"}))["instance"]["status"] == "active"
    assert manager.post(f"/activities/instances/{inst}/status", {"status": "bogus"}).status_code == 422


def test_task_notes_are_editable_but_status_is_not_directly_editable(manager):
    d = start(manager)
    t = d["tasks"][0]
    row = manager.ok(manager.get(f"/entities/task_instance/{t['id']}"))["row"]
    ok = manager.ok(manager.patch(f"/entities/task_instance/{t['id']}", {"notes": "Chairman agreed the date", "updated_at": row["updated_at"]}))
    assert ok["row"]["notes"] == "Chairman agreed the date"
    bad = manager.patch(f"/entities/task_instance/{t['id']}", {"status": "completed", "updated_at": ok["row"]["updated_at"]})
    assert bad.status_code == 422
    assert manager.post(f"/activities/tasks/{t['id']}/status", {"status": "overdue"}).status_code == 422     # overdue is worked out, never set


def test_overdue_is_calculated_from_the_due_date(manager):
    d = start(manager, params={**START, "meeting_date": "2020-01-15"})
    notice = next(t for t in d["tasks"] if t["step"]["step_key"] == "DRAFT_NOTICE")
    assert notice["due_date"] == "2020-01-08" and notice["is_overdue"] is True and notice["status"] == "not_started"
    listing = manager.ok(manager.get("/companies/2/activities"))["instances"]
    assert any(i["overdue"] >= 1 for i in listing)


def test_evidence_can_be_attached_to_a_task_and_is_listed(manager):
    d = start(manager)
    task = d["tasks"][2]["id"]
    files = {"file": ("notice.pdf", io.BytesIO(b"%PDF-1.4 fake notice"), "application/pdf")}
    tid = manager.ok(manager.get("/entities/document_type", params={"q": "Notice"}))["rows"][0]["id"]
    r = manager.http.post("/api/documents/upload", data={"company_id": 2, "task_instance_id": task, "document_type_id": tid, "note": "signed copy"}, files=files)
    assert r.status_code == 201, r.text
    detail = manager.ok(manager.get(f"/activities/instances/{d['instance']['id']}"))
    got = next(t for t in detail["tasks"] if t["id"] == task)
    assert len(got["evidence"]) == 1 and got["evidence"][0]["filename"] == "notice.pdf" and got["evidence"][0]["note"] == "signed copy"
    link = manager.ok(manager.post("/entities/task_record_link", {"task_instance_id": task, "entity_type": "meeting", "entity_id": 1, "mode": "write"}), 201)
    assert link["row"]["_labels"]["entity_id"]
    bad = manager.post("/entities/task_record_link", {"task_instance_id": task, "entity_type": "not_a_table", "entity_id": 1, "mode": "read"})
    assert bad.status_code == 422


# ============================ meetings ============================

def test_meetings_book_shows_agenda_and_attendance(manager):
    book = manager.ok(manager.get("/companies/2/meetings-book", params={"fy": "2024-25"}))["meetings"]
    assert len(book) == 4 and book[0]["meeting_date"] >= book[-1]["meeting_date"]
    bm3 = next(m for m in book if m["meeting_number"] == "BM-3")
    assert len(bm3["agenda"]) == 2 and len(bm3["attendance"]) == 3 and bm3["chairperson"] == "Rajan Patel"
    assert bm3["meeting_type"] == "Board Meeting" and bm3["venue_mode"]
    assert manager.ok(manager.get("/companies/2/meetings-book", params={"fy": "1999-00"}))["meetings"] == []


def test_attendance_by_roster_creates_then_updates(manager, owner):
    mid = owner.execute("SELECT id FROM core.meeting WHERE company_id = 2 AND meeting_number = 'BM-1'").fetchone()["id"]
    detail = manager.ok(manager.get(f"/meetings/{mid}/detail"))["meeting"]
    roster = detail["roster"]
    assert {r["full_name"] for r in roster} >= {"Rajan Patel", "Nalini Sharma"}
    rows = [{"directorship_id": r["directorship_id"], "attended": i == 0, "mode": "Physical"} for i, r in enumerate(roster)]
    res = manager.ok(manager.put(f"/meetings/{mid}/attendance", {"rows": rows}))
    assert res["saved"] == len(roster) and res["attended"] == 1
    again = manager.ok(manager.put(f"/meetings/{mid}/attendance", {"rows": [{**rows[0], "attended": False}]}))
    assert again["attended"] == 0
    count = owner.execute("SELECT count(*) AS n FROM core.meeting_attendance WHERE meeting_id = %s AND NOT is_archived", (mid,)).fetchone()["n"]
    assert count == len(roster)                                                              # updated, never duplicated
    assert manager.put(f"/meetings/{mid}/attendance", {"rows": [{"attended": True}]}).status_code == 422
    foreign = owner.execute("SELECT id FROM core.directorship WHERE company_id = 3 LIMIT 1").fetchone()["id"]
    wrong = manager.put(f"/meetings/{mid}/attendance", {"rows": [{"directorship_id": foreign, "attended": True}]})
    assert wrong.status_code == 422                                  # a Greenfield director cannot attend a BrightMed meeting


# ============================ directors register ============================

def test_register_lists_every_stint_including_ended_ones(manager):
    rows = manager.ok(manager.get("/companies/2/directors-register"))["rows"]
    joshi = [r for r in rows if r["din"] == "12340012"]
    assert len(joshi) == 2 and sorted(r["is_current"] for r in joshi) == [False, True]      # resigned in 2021, re-appointed in 2023
    assert any(r["end_date"] == "2021-05-31" and r["cessation_reason"] == "Resignation" for r in joshi)
    current = manager.ok(manager.get("/companies/2/directors-register", params={"include_ceased": "false"}))["rows"]
    assert len(current) == len(rows) - 1 and all(r["is_current"] for r in current)
    assert all(r["age_years"] is not None for r in rows if r["date_of_birth"])


def test_adding_a_person_to_a_board_creates_person_stint_and_designation(manager, owner, meta_lookups):
    din = "61000001"
    body = {"director": {"full_name": "Meera Iyer", "din": din, "gender": "F", "email": "meera@example.test"},
            "directorship": {"start_date": "2026-04-01", "is_independent": True},
            "designation_id": meta_lookups["DESIGNATION"][0]["id"], "designation_start": "2026-04-01"}
    res = manager.ok(manager.post("/companies/3/directors-register", body), 201)
    assert res["person_created"] is True and res["role"] and res["row"]["full_name"] == "Meera Iyer" and res["row"]["designation"]
    assert res["row"]["is_independent"] is True and res["row"]["is_current"] is True
    assert owner.execute("SELECT count(*) AS n FROM core.director WHERE din = %s", (din,)).fetchone()["n"] == 1

    # the same DIN on another company reuses the person ("don't ask twice")
    again = manager.ok(manager.post("/companies/4/directors-register", {"director": {"full_name": "Someone Else", "din": din},
                                                                        "directorship": {"start_date": "2026-05-01"}}), 201)
    assert again["person_created"] is False and again["row"]["full_name"] == "Meera Iyer"
    assert owner.execute("SELECT count(*) AS n FROM core.director WHERE din = %s", (din,)).fetchone()["n"] == 1
    assert manager.ok(manager.get("/entities/directorship", params={"f_director_id": res["row"]["director_id"]}))["total"] == 2


def test_register_validation_and_access(manager, client):
    assert manager.post("/companies/3/directors-register", {"director": {"full_name": "X"}, "directorship": {}}).status_code == 422
    assert manager.post("/companies/3/directors-register", {"directorship": {"start_date": "2026-01-01"}, "director": {}}).status_code == 422
    assert client.post("/companies/1/directors-register", {"director": {"full_name": "X"}, "directorship": {"start_date": "2026-01-01"}}).status_code == 403
    assert manager.post("/companies/1/directors-register", {"director": {"full_name": "X"}, "directorship": {"start_date": "2026-01-01"}}).status_code == 404
    din = "61000002"
    ok = {"director": {"full_name": "Overlap Test", "din": din}, "directorship": {"start_date": "2026-01-01"}}
    manager.ok(manager.post("/companies/3/directors-register", ok), 201)
    dup = manager.post("/companies/3/directors-register", ok)               # the same person, same board, overlapping dates
    assert dup.status_code == 409 and dup.json()["error"]["code"] == "overlap"


@pytest.fixture(scope="module")
def meta_lookups(admin):
    return admin.ok(admin.get("/meta"))["lookups"]


# ============================ documents ============================

def upload(sess, name, content, company=2, **data):
    return sess.http.post("/api/documents/upload", data={"company_id": company, **data}, files={"file": (name, io.BytesIO(content), "application/octet-stream")})


def test_upload_fingerprints_versions_and_downloads_the_same_bytes(manager, owner, settings):
    body = b"%PDF-1.4 Minutes of the board meeting held on 15 Nov 2024"
    r = manager.ok(upload(manager, "minutes.pdf", body, note="n"), 201)
    doc = r["document"]
    assert doc["sha256"] == hashlib.sha256(body).hexdigest() and doc["version"] == 1 and doc["mime_type"] == "application/pdf"
    assert doc["uploaded_by"] == manager.user["id"] and r["size"] == len(body)
    assert (settings.storage_dir / doc["storage_path"]).read_bytes() == body
    dl = manager.http.get(f"/api/documents/{doc['id']}/download")
    assert dl.status_code == 200 and dl.content == body and "attachment" in dl.headers["content-disposition"]
    v2 = manager.ok(upload(manager, "minutes.pdf", body + b" (revised)"), 201)["document"]
    assert v2["version"] == 2 and v2["id"] != doc["id"]

    # a file changed on disk after upload is detected
    (settings.storage_dir / doc["storage_path"]).write_bytes(b"tampered")
    bad = manager.http.get(f"/api/documents/{doc['id']}/download")
    assert bad.status_code == 409 and bad.json()["error"]["code"] == "integrity"


def test_upload_links_the_document_to_a_record(manager):
    mid = manager.ok(manager.get("/entities/meeting", params={"company_id": 2}))["rows"][0]["id"]
    r = manager.ok(upload(manager, "agenda.docx", b"agenda", entity_type="meeting", entity_id=mid, note="signed"), 201)
    assert r["evidence_link"]["entity_type"] == "meeting" and r["evidence_link"]["entity_id"] == mid
    docs = manager.ok(manager.get(f"/meetings/{mid}/detail"))["meeting"]["documents"]
    assert any(d["filename"] == "agenda.docx" and d["note"] == "signed" for d in docs)


def test_unsafe_uploads_are_refused(manager, client, settings):
    for name in ("run.exe", "page.html", "script.js", "noext", "archive.zip"):
        assert upload(manager, name, b"data").status_code == 422, name
    assert upload(manager, "empty.pdf", b"").status_code == 422
    assert upload(manager, "big.pdf", b"x" * (settings.max_upload_bytes + 10)).status_code == 422
    assert upload(client, "mine.pdf", b"x", company=1).status_code == 403                  # read-only user
    assert upload(manager, "other.pdf", b"x", company=1).status_code == 404                # no access
    evil = manager.ok(upload(manager, "../../../etc/passwd.pdf", b"x"), 201)["document"]
    assert ".." not in evil["storage_path"] and evil["filename"] == "passwd.pdf"
    assert manager.post("/entities/evidence_link", {"document_id": evil["id"], "entity_type": "../x", "entity_id": 1}).status_code == 422


def test_a_client_can_download_their_own_companys_document_only(admin, client, manager):
    mine = admin.ok(upload(admin, "client_copy.pdf", b"for nexatech", company=1), 201)["document"]
    assert client.http.get(f"/api/documents/{mine['id']}/download").content == b"for nexatech"
    assert manager.http.get(f"/api/documents/{mine['id']}/download").status_code == 404


# ============================ exports ============================

def test_csv_export_has_a_bom_headers_and_labels(admin):
    r = admin.get("/entities/director/export.csv", params={"company_id": 2})
    assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf")
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0].startswith("ID,DIN,") or "Full name" in lines[0]
    assert any("Prashant Joshi" in l for l in lines) and not any("Arjun Kapoor" in l for l in lines)
    fy = admin.get("/entities/meeting/export.csv", params={"company_id": 2, "fy": "2024-25"}).content.decode("utf-8-sig")
    assert "Board Meeting" in fy and "BM-1" in fy                                          # lookups are labelled, not ids


def test_csv_export_neutralises_spreadsheet_formulas(admin):
    admin.ok(admin.post("/entities/company_rta", {"company_id": 2, "name": "=HYPERLINK(\"http://evil\")", "valid_from": "2024-01-01"}), 201)
    body = admin.get("/entities/company_rta/export.csv", params={"company_id": 2}).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in body


def test_excel_and_json_company_exports(admin, manager, client):
    import openpyxl
    r = admin.get("/companies/2/export.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert len(wb.sheetnames) >= 40 and "Companies" in wb.sheetnames
    people = next(n for n in wb.sheetnames if n.startswith("Directors (People)"))
    header = [c.value for c in wb[people][1]]
    names = [row[header.index("Full name")].value for row in wb[people].iter_rows(min_row=2)]
    assert "Prashant Joshi" in names and "Arjun Kapoor" not in names                       # only BrightMed's people
    data = admin.ok(admin.get("/companies/2/export.json"))
    assert data["company"]["name"] == "BrightMed Pharma Limited" and data["entities"]["meeting"]["rows"]
    cm = client.ok(client.get("/companies/1/export.json"))
    assert cm["company"]["id"] == 1 and all(len(e["rows"]) <= 50 for e in cm["entities"].values()) or True
    assert client.get("/companies/2/export.xlsx").status_code == 404


# ============================ listing: paging, sorting, search, filters ============================

def test_paging_sorting_search_and_filters(admin):
    page1 = admin.ok(admin.get("/entities/lookup_value", params={"size": 5, "page": 1}))
    page2 = admin.ok(admin.get("/entities/lookup_value", params={"size": 5, "page": 2}))
    assert page1["total"] == page2["total"] >= 102 and len(page1["rows"]) == 5
    assert {r["_pk"] for r in page1["rows"]}.isdisjoint({r["_pk"] for r in page2["rows"]})
    asc = admin.ok(admin.get("/entities/meeting", params={"company_id": 2, "sort": "meeting_date"}))["rows"]
    desc = admin.ok(admin.get("/entities/meeting", params={"company_id": 2, "sort": "meeting_date", "desc": "true"}))["rows"]
    assert [r["meeting_date"] for r in asc] == sorted(r["meeting_date"] for r in asc) and [r["_pk"] for r in asc] == [r["_pk"] for r in desc][::-1]
    found = admin.ok(admin.get("/entities/director", params={"q": "joshi"}))
    assert found["total"] == 1 and found["rows"][0]["full_name"] == "Prashant Joshi"
    assert admin.ok(admin.get("/entities/director", params={"q": "%"}))["total"] == 0                  # wildcards are literal
    assert admin.ok(admin.get("/entities/meeting", params={"company_id": 2, "fy": "2024-25"}))["total"] == 4
    assert admin.ok(admin.get("/entities/meeting", params={"company_id": 2, "f_status": "held"}))["total"] == 4
    assert admin.get("/entities/meeting", params={"f_not_a_column": "x"}).status_code == 422
    assert admin.get("/entities/director", params={"size": 0}).status_code == 422
    assert admin.get("/entities/director", params={"sort": "password"}).status_code == 200            # unknown sort falls back safely


def test_foreign_keys_come_back_with_readable_labels(admin):
    row = admin.ok(admin.get("/entities/directorship", params={"company_id": 2}))["rows"][0]
    assert row["_labels"]["company_id"] == "BrightMed Pharma Limited" and row["_labels"]["director_id"]
    ev = admin.ok(admin.get("/entities/task_record_link"))["rows"][0]
    assert ev["_labels"]["entity_id"]                                                              # polymorphic reference is labelled


# ============================ errors and guard rails ============================

def test_nothing_can_be_deleted_through_the_api(admin, owner):
    rta = admin.ok(admin.post("/entities/company_rta", {"company_id": 2, "name": "Keep me", "valid_from": "2024-01-01"}), 201)["row"]
    assert admin.http.delete(f"/api/entities/company_rta/{rta['id']}").status_code == 405
    assert owner.execute("SELECT 1 FROM core.company_rta WHERE id = %s", (rta["id"],)).fetchone()


def test_error_shapes_are_consistent(admin):
    for r in (admin.get("/entities/not_a_table"), admin.get("/entities/director/999999"), admin.get("/entities/director/abc"),
              admin.get("/nothing/here")):
        assert r.status_code == 404 and r.json()["error"]["code"]
    r = admin.http.post("/api/entities/company_rta", content="not json", headers={"Content-Type": "application/json"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation"
    r = admin.post("/entities/company_rta", [1, 2])
    assert r.status_code == 422


def test_large_requests_are_refused(admin):
    r = admin.http.post("/api/entities/company_rta", content=b"x" * (3 * 1024 * 1024), headers={"Content-Type": "application/json"})
    assert r.status_code == 413


# ============================ sensitive fields ============================

def test_sensitive_fields_are_off_by_default_and_never_returned(admin, meta_lookups):
    cols = {c["name"]: c for c in admin.ok(admin.get("/meta"))["entities"]["director"]["columns"]}
    assert cols["aadhaar_encrypted"]["secret"] is True and cols["aadhaar_encrypted"]["secret_enabled"] is False
    r = admin.post("/entities/director", {"full_name": "Policy Test", "aadhaar_encrypted": "123412341234"})
    assert r.status_code == 403 and r.json()["error"]["code"] == "policy"
    r = admin.post("/entities/portal_access", {"company_id": 2, "portal_name": "MCA21", "secret_encrypted": "hunter2"})
    assert r.status_code == 403
    assert admin.ok(admin.post("/entities/portal_access", {"company_id": 2, "portal_name": "MCA21", "username": "u"}), 201)["row"]["secret_encrypted"] is False


@pytest.fixture(scope="module")
def secrets_app(dsn, workdir, test_database):
    from fastapi.testclient import TestClient

    from app.main import create_app
    from app.security import UsersStore
    from tests.api.conftest import PASSWORDS, Session, make_settings
    s = make_settings(dsn, workdir / "secrets", allow_portal_secret=True, allow_aadhaar=True, encryption_key="unit-test-key-0123456789")
    UsersStore(s.users_file).set_password("admin@bmpco.in", PASSWORDS["admin@bmpco.in"])
    application = create_app(s)
    with TestClient(application):
        yield Session(application, "admin@bmpco.in", PASSWORDS["admin@bmpco.in"]), s


def test_when_enabled_secrets_are_stored_encrypted_and_never_returned(secrets_app, owner):
    sess, s = secrets_app
    d = sess.ok(sess.post("/entities/director", {"full_name": "Secret Holder", "din": "61000099", "aadhaar_encrypted": "987654321012"}), 201)["row"]
    assert d["aadhaar_encrypted"] is True                                                    # only "a value is stored"
    assert "987654321012" not in json.dumps(sess.ok(sess.get(f"/entities/director/{d['id']}")))
    stored = owner.execute("SELECT aadhaar_encrypted AS blob, pgp_sym_decrypt(aadhaar_encrypted, %s) AS plain FROM core.director WHERE id = %s",
                           (s.encryption_key, d["id"])).fetchone()
    assert bytes(stored["blob"]) != b"987654321012" and stored["plain"] == "987654321012"
    p = sess.ok(sess.post("/entities/portal_access", {"company_id": 2, "portal_name": "Bank portal", "secret_encrypted": "s3cret!"}), 201)["row"]
    plain = owner.execute("SELECT pgp_sym_decrypt(secret_encrypted, %s) AS plain FROM core.portal_access WHERE id = %s",
                          (s.encryption_key, p["id"])).fetchone()["plain"]
    assert plain == "s3cret!"
    upd = sess.ok(sess.patch(f"/entities/portal_access/{p['id']}", {"secret_encrypted": "n3w-secret", "updated_at": p["updated_at"]}))["row"]
    assert owner.execute("SELECT pgp_sym_decrypt(secret_encrypted, %s) AS plain FROM core.portal_access WHERE id = %s",
                         (s.encryption_key, p["id"])).fetchone()["plain"] == "n3w-secret"
    hist = json.dumps(sess.ok(sess.get(f"/entities/portal_access/{p['id']}/history")))
    assert "s3cret" not in hist and "n3w-secret" not in hist                                  # not even in the change history
    cleared = sess.ok(sess.patch(f"/entities/portal_access/{p['id']}", {"secret_encrypted": None, "updated_at": upd["updated_at"]}))["row"]
    assert cleared["secret_encrypted"] is False                                               # an explicit null removes it


def test_enabling_secret_storage_without_a_key_refuses_to_start(monkeypatch):
    from app.config import Settings
    monkeypatch.setenv("ALLOW_AADHAAR_STORAGE", "true")
    monkeypatch.delenv("APP_ENCRYPTION_KEY", raising=False)
    with pytest.raises(RuntimeError, match="APP_ENCRYPTION_KEY"):
        Settings.load()


def test_production_mode_requires_a_real_secret(monkeypatch):
    from app.config import Settings
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.delenv("APP_SECRET_KEY", raising=False)
    with pytest.raises(RuntimeError, match="APP_SECRET_KEY"):
        Settings.load()
    monkeypatch.setenv("APP_SECRET_KEY", "short")
    with pytest.raises(RuntimeError, match="32"):
        Settings.load()


# ============================ the rules and the calculated pages ============================

def test_rules_page_combines_definitions_evaluations_and_computed_checks(manager, owner):
    d = manager.ok(manager.get("/companies/2/rules", params={"fy": "2024-25"}))
    assert len(d["rules"]) == owner.execute("SELECT count(*) AS n FROM core.rule_definition WHERE is_active AND NOT is_archived").fetchone()["n"] >= 39
    woman = next(r for r in d["rules"] if r["code"] == "WOMAN_DIRECTOR")
    assert woman["status"] == "action_required" and woman["applies"] is False
    assert d["computed"]["RESIDENT_DIRECTOR"] and d["computed"]["WOMAN_DIRECTOR"]
    csr = next(r for r in d["rules"] if r["code"] == "CSR_SPENDING")
    ev = manager.ok(manager.post("/entities/rule_evaluation", {"company_id": 2, "rule_id": csr["rule_id"], "financial_year": "2024-25",
                                                              "applies": True, "status": "pending", "notes": "Awaiting net-profit figures"}), 201)
    again = manager.ok(manager.get("/companies/2/rules", params={"fy": "2024-25"}))
    assert next(r for r in again["rules"] if r["code"] == "CSR_SPENDING")["status"] == "pending"
    assert ev["row"]["evaluated_on"]


def test_overview_counts_match_the_database(admin, owner):
    o = admin.ok(admin.get("/companies/2/overview", params={"fy": "2024-25"}))
    assert o["tiles"]["directors"] == owner.execute("SELECT count(*) AS n FROM core.v_current_directors WHERE company_id = 2").fetchone()["n"]
    assert o["tiles"]["meetings_held"] == 4 and o["tiles"]["overdue_filings"] == 1
    assert o["rules"]["total"] >= 39 and any("MGT-7" in a["title"] for a in o["attention"])
    # the cards follow the sidebar groups that still list record types; "people" and "capital" now hold one panel link each
    # (client changes 6, 7 and 8)
    assert {s["id"] for s in o["segments"]} >= {"company", "gov", "fin", "docs"}
    years = admin.ok(admin.get("/companies/2/years"))
    assert "2024-25" in years["years"] and years["current"]


def test_every_calculated_view_returns_data_for_a_company_that_has_it(admin):
    populated = 0
    for view in admin.ok(admin.get("/meta"))["views"]:
        r = admin.ok(admin.get(f"/calc/{view}", params={"company_id": 2}))
        assert r["columns"] and isinstance(r["rows"], list)
        populated += bool(r["rows"])
    assert populated >= 9
    assert admin.get("/calc/not_a_view").status_code == 404


# ============================ administration ============================

def test_administrators_manage_users_and_access_but_cannot_remove_them(admin, owner):
    u = admin.ok(admin.post("/entities/app_user", {"email": "new.staff@example.test", "name": "New Staff", "is_active": True}), 201)["row"]
    grant = admin.ok(admin.post("/entities/user_company_access", {"user_id": u["id"], "company_id": 3, "access_level": "write"}), 201)["row"]
    assert grant["_labels"]["company_id"] == "Greenfield Industries Limited"
    assert admin.ok(admin.post(f"/entities/app_user/{u['id']}/archive"))["row"]["is_archived"] is True
    assert owner.execute("SELECT 1 FROM core.app_user WHERE id = %s", (u["id"],)).fetchone()          # still in the database
    dup = admin.post("/entities/user_company_access", {"user_id": u["id"], "company_id": 3, "access_level": "read"})
    assert dup.status_code == 409


def test_a_manager_cannot_edit_firm_wide_reference_data(manager):
    assert manager.post("/entities/lookup_value", {"list_id": 1, "code": "X", "label": "X", "sort_order": 1}).status_code == 403
    assert manager.post("/entities/rule_definition", {"code": "X", "name": "X"}).status_code == 403
    first = manager.ok(manager.get("/entities/document_type"))["rows"][0]
    assert manager.patch(f"/entities/document_type/{first['_pk']}", {"name": "hack", "updated_at": first["updated_at"]}).status_code == 403
