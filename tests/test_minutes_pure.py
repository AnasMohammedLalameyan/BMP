"""Compliance engine and Excel builder - pure tests, no database."""
import io

from openpyxl import load_workbook

from app.minutes_compliance import compute
from app.minutes_excel import HEADERS, WIDTHS, build_meetings_workbook


def chips(register, meetings):
    return {c["key"]: c for c in compute(register, meetings)}


def m(date, **kw):
    return {"meeting_date": date, **kw}


def test_empty_register_only_min_chip():
    assert [c["key"] for c in compute("board", [])] == ["min"]
    assert compute("egm", []) == []


def test_count_shortfall_does_not_flip_other_checks():
    c = chips("board", [m("2025-04-15", notice_date="2025-04-01")])
    assert c["min"]["status"] == "attention"
    assert c["notice"]["status"] == "ok"
    assert c["quorum"]["status"] == "unknown" and c["minutes"]["status"] == "unknown"
    assert "draft" not in c and "gap" not in c


def test_short_notice_ss1_and_ss2():
    c = chips("board", [m("2025-04-15", notice_date="2025-04-10"), m("2025-08-01", notice_date="2025-07-20")])
    assert c["notice"]["status"] == "attention" and c["notice"]["detail"] == "1 of 2 short notice"
    # SS-2: 21 CLEAR days -> 22 days between notice and meeting
    ok = chips("agm", [m("2025-09-30", notice_date="2025-09-08")])
    bad = chips("agm", [m("2025-09-30", notice_date="2025-09-10")])
    assert ok["notice"]["status"] == "ok" and bad["notice"]["status"] == "attention"


def test_gap_over_120_days():
    c = chips("board", [m("2025-04-01"), m("2025-08-05")])
    assert c["gap"]["status"] == "attention" and "126" in c["gap"]["detail"]
    assert chips("board", [m("2025-04-01"), m("2025-07-01")])["gap"]["status"] == "ok"
    assert "gap" not in chips("nrc", [m("2025-04-01"), m("2025-09-05")])


def att(n_present, n_total):
    return [{"attended": i < n_present} for i in range(n_total)]


def test_quorum():
    assert chips("board", [m("2025-04-01", attendance=att(1, 6))])["quorum"]["status"] == "attention"
    assert chips("board", [m("2025-04-01", attendance=att(2, 6))])["quorum"]["status"] == "ok"
    assert chips("board", [m("2025-04-01", attendance=att(3, 9))])["quorum"]["status"] == "ok"
    # typed numbers are used only when nothing is ticked
    assert chips("board", [m("2025-04-01", attendance_present=1, attendance_total=3)])["quorum"]["status"] == "attention"


def test_minutes_signed_late_and_draft():
    c = chips("board", [m("2025-04-01", signing_date="2025-05-10", draft_circulated="2025-04-20")])
    assert c["minutes"]["status"] == "attention" and c["draft"]["status"] == "attention"
    c = chips("board", [m("2025-04-01", signing_date="2025-04-30", draft_circulated="2025-04-10")])
    assert c["minutes"]["status"] == "ok" and c["draft"]["status"] == "ok"


def test_director_present_outside_directorship_dates():
    c = chips("board", [m("2025-04-01", attendance=att(3, 3), outside_directorship=["X"])])
    assert c["board_register"]["status"] == "attention"
    assert chips("board", [m("2025-04-01", attendance=att(3, 3))])["board_register"]["status"] == "ok"


def test_excel_layout():
    data = {"company_name": "Acme Ltd", "financial_year": "2025-26",
            "registers": [{"key": "board", "name": "Board Meetings"}, {"key": "audit", "name": "Audit Committee"}],
            "meetings": [{"meeting_id": 2, "register": "board", "meeting_date": "2025-08-05", "notice_date": "2025-07-20",
                          "agenda": [{"text": "A", "forms": "MGT-14, DIR-12", "docs": "Form DIR-8"},
                                     {"text": "", "forms": "", "docs": ""},
                                     {"text": "B", "forms": "MGT-14", "docs": "Form DIR-8, X"}],
                          "directors_present": "A, B", "comments": "=evil"},
                         {"meeting_id": 1, "register": "board", "meeting_date": "2025-04-01", "agenda": []}]}
    wb = load_workbook(io.BytesIO(build_meetings_workbook(data)))
    assert wb.sheetnames == ["BM"]
    ws = wb["BM"]
    assert ws["A1"].value == "Acme Ltd — Board Meetings — FY 2025-26"
    assert [c.value for c in ws[2]] == HEADERS
    assert [round(ws.column_dimensions[c.column_letter].width) for c in ws[2]] == WIDTHS
    assert ws["C3"].value.day == 1 and ws["C4"].number_format == "dd.mm.yyyy"     # sorted by date
    assert ws["I4"].value == "1. A\n2. B"
    assert ws["J4"].value == "MGT-14\nDIR-12" and ws["K4"].value == "Form DIR-8\nX"
    assert ws["O4"].value == "'=evil"
    assert ws.freeze_panes == "A3" and ws.auto_filter.ref.startswith("A2:")


def test_excel_empty_state():
    wb = load_workbook(io.BytesIO(build_meetings_workbook(
        {"company_name": "Acme", "financial_year": "2025-26", "registers": [], "meetings": []})))
    assert wb.sheetnames == ["Meetings"] and wb["Meetings"]["A2"].value == "No meetings recorded for this year."
