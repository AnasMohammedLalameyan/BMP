"""Excel workbook for the meetings of one company and financial year. Pure: no database, no network."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SHEET = {"board": "BM", "audit": "ACM", "nrc": "NRCM", "src": "SRCM", "csr": "CSRM", "rmc": "RMCM",
         "agm": "AGM", "egm": "EGM"}
HEADERS = ["Serial No.", "Date of Notice", "Date of Meeting", "Time", "Chairman", "Members present",
           "Leave of absence", "Attendance Sheet", "Minutes (agenda)", "Forms filed", "Documents to be checked",
           "Date of Entry", "Date of Signing", "Place of Signing", "Remarks"]
WIDTHS = [10, 14, 14, 20, 20, 26, 16, 14, 52, 20, 30, 13, 13, 15, 24]
NAVY = "0F2747"
_side = Side(style="thin", color="D9DEE8")
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)
DATE_FMT = "dd.mm.yyyy"


def _date(s):
    try:
        return date.fromisoformat(str(s)[:10]) if s else None
    except ValueError:
        return None


def _unique(agenda, field):
    seen, out = set(), []
    for it in agenda or []:
        for part in str(it.get(field) or "").split(","):
            part = part.strip()
            if part and part not in seen:
                seen.add(part)
                out.append(part)
    return "\n".join(out)


def _safe(v):
    """Stop a spreadsheet reading text a person typed as a formula."""
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def build_meetings_workbook(data: dict) -> bytes:
    company, fy = data.get("company_name", ""), data.get("financial_year", "")
    wb = Workbook()
    wb.remove(wb.active)
    any_sheet = False
    for reg in data.get("registers", []):
        rows = sorted((m for m in data.get("meetings", []) if m.get("register") == reg["key"]),
                      key=lambda m: (m.get("meeting_date") or "", m.get("meeting_id") or 0))
        if not rows:
            continue
        any_sheet = True
        ws = wb.create_sheet(SHEET.get(reg["key"], reg["key"].upper())[:31])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
        ws["A1"] = f"{company} — {reg['name']} — FY {fy}"
        ws["A1"].font = Font(bold=True, color=NAVY, size=14)
        for i, h in enumerate(HEADERS, 1):
            c = ws.cell(row=2, column=i, value=h)
            c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = BORDER
        ws.row_dimensions[2].height = 26
        for i, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = w
        for n, m in enumerate(rows, 1):
            agenda = m.get("agenda") or []
            items = [a for a in agenda if (a.get("text") or "").strip()]
            minutes = "\n".join(f"{k}. {it['text']}" for k, it in enumerate(items, 1))
            values = [n, _date(m.get("notice_date")), _date(m.get("meeting_date")), m.get("meeting_time"),
                      m.get("chairman"), m.get("directors_present"), m.get("leave_of_absence"),
                      m.get("attendance_sheet"), minutes, _unique(agenda, "forms"), _unique(agenda, "docs"),
                      _date(m.get("entry_date")), _date(m.get("signing_date")), m.get("signing_place"),
                      m.get("comments")]
            for i, v in enumerate(values, 1):
                c = ws.cell(row=2 + n, column=i, value=_safe(v))
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = BORDER
                if isinstance(v, date):
                    c.number_format = DATE_FMT
        ws.freeze_panes = "A3"
        last = ws.cell(row=2, column=len(HEADERS)).column_letter
        ws.auto_filter.ref = f"A2:{last}{2 + len(rows)}"
    if not any_sheet:
        ws = wb.create_sheet("Meetings")
        ws["A1"] = f"{company} — Minutes & Meetings — FY {fy}"
        ws["A1"].font = Font(bold=True, color=NAVY, size=14)
        ws["A2"] = "No meetings recorded for this year."
        ws["A2"].font = Font(italic=True, color="808080")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
