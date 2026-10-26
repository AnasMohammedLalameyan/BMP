"""Companies: the landing list, onboarding a new company, the overview page, years, exports."""
from __future__ import annotations

import io
import json
from dataclasses import replace
from datetime import date, timedelta

from fastapi import APIRouter, Body, Depends, Request
from starlette.responses import Response

from ..access import UserCtx
from ..deps import admin_user, current_user, privileged_session, readable_company, session
from ..errors import not_found
from ..registry import ENTITIES, NAV, SCOPE
from ..responses import ApiResponse, ok
from .meta import humanize

router = APIRouter(tags=["companies"])


def current_financial_year(today: date | None = None) -> str:
    today = today or date.today()
    start = today.year if today.month >= 4 else today.year - 1
    return f"{start}-{(start + 1) % 100:02d}"


@router.get("/companies")
def list_companies(request: Request, ctx: UserCtx = Depends(current_user)):
    with session(request, ctx, readonly=True) as conn:
        rows = conn.execute(
            """SELECT c.id, c.name, c.cin, c.pan, c.is_listed, c.status, c.date_of_incorporation,
                      (SELECT lv.label FROM core.lookup_value lv WHERE lv.id = c.category_id) AS category,
                      (SELECT lv.label FROM core.lookup_value lv WHERE lv.id = c.subcategory_id) AS subcategory,
                      (SELECT count(*) FROM core.v_current_directors v WHERE v.company_id = c.id) AS directors,
                      (SELECT count(*) FROM core.v_filings_overdue v WHERE v.company_id = c.id) AS overdue_filings
               FROM core.company c
               WHERE c.id = ANY(%(ids)s::bigint[]) AND NOT c.is_archived
               ORDER BY c.name""", {"ids": ctx.readable}).fetchall()
    for r in rows:
        r["access_level"] = ctx.level_for(r["id"])
    return ok({"companies": rows})


@router.post("/companies")
def create_company(request: Request, payload: dict = Body(...), ctx: UserCtx = Depends(admin_user)):
    """Onboard a company. Row-level security means the application login cannot insert a company that
    nobody has access to yet, so this one action runs on the privileged connection, and in the same
    transaction grants the creating administrator access to it."""
    with privileged_session(request, ctx) as conn:
        priv = replace(ctx, privileged=True)
        row = request.app.state.engine.create(conn, "company", priv, payload)
        conn.execute("INSERT INTO core.user_company_access (user_id, company_id, access_level) VALUES (%s, %s, 'admin')",
                     (ctx.id, row["id"]))
    return ok({"row": row}, 201)


def _company_row(conn, company_id: int) -> dict:
    row = conn.execute(
        """SELECT c.id, c.name, c.cin, c.is_listed, c.status,
                  (SELECT lv.label FROM core.lookup_value lv WHERE lv.id = c.category_id) AS category
           FROM core.company c WHERE c.id = %s""", (company_id,)).fetchone()
    if row is None:
        raise not_found("Company")
    return row


@router.get("/companies/{company_id}/years")
def financial_years(company_id: int, request: Request, ctx: UserCtx = Depends(current_user)):
    readable_company(ctx, company_id)
    catalog = request.app.state.catalog
    parts = [f'SELECT "financial_year" AS fy FROM {t.ident} WHERE "company_id" = %(cid)s AND "financial_year" IS NOT NULL'
             for t in catalog.base_tables() if t.has("company_id") and t.has("financial_year")]
    with session(request, ctx, readonly=True) as conn:
        rows = conn.execute(" UNION ".join(parts), {"cid": company_id}).fetchall() if parts else []
    # "years" = years that have records; "current" = the year we are in now (offered even if empty)
    return ok({"years": sorted({r["fy"] for r in rows}, reverse=True), "current": current_financial_year()})


@router.get("/companies/{company_id}/overview")
def overview(company_id: int, request: Request, fy: str | None = None, ctx: UserCtx = Depends(current_user)):
    readable_company(ctx, company_id)
    engine = request.app.state.engine
    fy = fy or current_financial_year()
    today = date.today()
    with session(request, ctx, readonly=True) as conn:
        company = _company_row(conn, company_id)
        one = lambda sql, **p: conn.execute(sql, {"cid": company_id, "fy": fy, **p}).fetchone()
        directors = one("""SELECT count(*) AS total, count(*) FILTER (WHERE is_independent) AS independent,
                                  count(*) FILTER (WHERE gender = 'F') AS women
                           FROM core.v_current_directors WHERE company_id = %(cid)s""")
        holders = one("SELECT coalesce(max(current_shareholder_count), 0) AS n FROM core.v_company_shareholder_counts "
                      "WHERE company_id = %(cid)s")["n"]
        meetings = one("""SELECT count(*) FILTER (WHERE status = 'held') AS held,
                                 count(*) FILTER (WHERE status = 'planned') AS planned
                          FROM core.meeting WHERE company_id = %(cid)s AND financial_year = %(fy)s AND NOT is_archived""")
        filings = one("""SELECT count(*) AS total, count(*) FILTER (WHERE filing_date IS NOT NULL) AS filed
                         FROM core.government_filing WHERE company_id = %(cid)s AND NOT is_archived""")
        overdue = conn.execute("SELECT form_name, financial_year, due_date, days_overdue FROM core.v_filings_overdue "
                               "WHERE company_id = %(cid)s ORDER BY due_date", {"cid": company_id}).fetchall()
        rule_rows = conn.execute(
            """SELECT r.id, r.code, r.name, e.status, e.applies, e.evaluated_on
               FROM core.rule_definition r
               LEFT JOIN LATERAL (SELECT status, applies, evaluated_on FROM core.rule_evaluation x
                                  WHERE x.rule_id = r.id AND x.company_id = %(cid)s AND x.financial_year = %(fy)s
                                    AND NOT x.is_archived ORDER BY x.evaluated_on DESC LIMIT 1) e ON TRUE
               WHERE r.is_active AND NOT r.is_archived""", {"cid": company_id, "fy": fy}).fetchall()
        rules = {"met": 0, "pending": 0, "action_required": 0, "not_applicable": 0, "unevaluated": 0}
        action_rules = []
        for r in rule_rows:
            key = r["status"] or "unevaluated"
            rules[key] = rules.get(key, 0) + 1
            if r["status"] == "action_required":
                action_rules.append({"code": r["code"], "name": r["name"]})
        rules["total"] = len(rule_rows)
        tasks = one("""SELECT count(*) FILTER (WHERE t.status NOT IN ('completed', 'skipped')) AS open,
                              count(*) FILTER (WHERE t.status NOT IN ('completed', 'skipped')
                                               AND t.due_date < current_date) AS overdue
                       FROM ops.task_instance t JOIN ops.activity_instance i ON i.id = t.instance_id
                       WHERE i.company_id = %(cid)s AND NOT t.is_archived AND NOT i.is_archived""")
        dsc = conn.execute(
            """SELECT d.full_name, d.dsc_valid_until FROM core.directorship ds JOIN core.director d ON d.id = ds.director_id
               WHERE ds.company_id = %(cid)s AND NOT ds.is_archived AND (ds.end_date IS NULL OR ds.end_date >= current_date)
                 AND d.dsc_valid_until IS NOT NULL AND d.dsc_valid_until <= current_date + 90
               ORDER BY d.dsc_valid_until""", {"cid": company_id}).fetchall()
        upcoming = conn.execute(
            """SELECT id, meeting_number, meeting_date FROM core.meeting
               WHERE company_id = %(cid)s AND status = 'planned' AND NOT is_archived ORDER BY meeting_date LIMIT 5""",
            {"cid": company_id}).fetchall()

        attention = []
        for f in overdue:
            attention.append({"level": "no", "title": f"{f['form_name']} overdue", "detail":
                              f"Due {f['due_date']} - {f['days_overdue']} days overdue", "view": "entity:government_filing"})
        for r in action_rules:
            attention.append({"level": "no", "title": f"Rule needs action: {r['name']}", "detail": r["code"], "view": "rules"})
        for d in dsc:
            expired = d["dsc_valid_until"] < today
            attention.append({"level": "no" if expired else "part",
                              "title": f"DSC {'expired' if expired else 'expiring'}: {d['full_name']}",
                              "detail": f"Valid until {d['dsc_valid_until']}", "view": "directors"})
        if tasks["overdue"]:
            attention.append({"level": "part", "title": f"{tasks['overdue']} task(s) past their due date",
                              "detail": "Activities & Tasks", "view": "activities"})
        for m in upcoming:
            attention.append({"level": "prog", "title": f"Planned meeting {m['meeting_number'] or ''}".strip(),
                              "detail": f"On {m['meeting_date']}", "view": "meetings"})

        segments = []
        for group in NAV:
            ents = []
            for item in group["items"]:
                key = item.get("entity")
                if not key or key not in ENTITIES:
                    continue
                if SCOPE.get(key, {"type": "direct"})["type"] not in ("direct", "parent", "masters"):
                    continue
                ents.append({"key": key, "label": ENTITIES[key]["label"],
                             "count": engine.count(conn, key, ctx, company_id=company_id)})
            if ents:
                segments.append({"id": group["id"], "title": group["title"], "icon": group["icon"], "entities": ents})

    return ok({
        "company": company, "fy": fy,
        "tiles": {"directors": directors["total"], "independent": directors["independent"], "women": directors["women"],
                  "shareholders": holders, "meetings_held": meetings["held"], "meetings_planned": meetings["planned"],
                  "filings": filings["total"], "filed": filings["filed"], "overdue_filings": len(overdue),
                  "open_tasks": tasks["open"], "overdue_tasks": tasks["overdue"]},
        "rules": rules, "attention": attention, "segments": segments,
    })


def _export_payload(request: Request, ctx: UserCtx, company_id: int) -> dict:
    engine = request.app.state.engine
    out: dict = {}
    with session(request, ctx, readonly=True) as conn:
        company = _company_row(conn, company_id)
        for key, cfg in ENTITIES.items():
            t = request.app.state.catalog.get(key)
            if t is None or SCOPE.get(key, {"type": "direct"})["type"] not in ("direct", "parent", "masters"):
                continue
            data = engine.list(conn, key, ctx, company_id=company_id, page=1, size=5000)
            cols = [c for c in t.columns if c.kind not in ("bytea", "vector")]
            out[key] = {"label": cfg["label"], "columns": [c.name for c in cols],
                        "headers": [humanize(c.name) for c in cols],
                        "rows": [[r["_labels"].get(c.name, r.get(c.name)) for c in cols] for r in data["rows"]]}
    return {"company": company, "exported_on": date.today().isoformat(), "entities": out}


@router.get("/companies/{company_id}/export.json")
def export_json(company_id: int, request: Request, ctx: UserCtx = Depends(current_user)):
    readable_company(ctx, company_id)
    payload = _export_payload(request, ctx, company_id)
    return ApiResponse(payload, headers={"Content-Disposition": f'attachment; filename="company_{company_id}.json"'})


@router.get("/companies/{company_id}/export.xlsx")
def export_xlsx(company_id: int, request: Request, ctx: UserCtx = Depends(current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    readable_company(ctx, company_id)
    payload = _export_payload(request, ctx, company_id)
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for key, e in payload["entities"].items():
        title = "".join(ch for ch in e["label"] if ch not in '[]:*?/\\')[:31] or key[:31]
        while title in used:
            title = title[:28] + str(len(used))
        used.add(title)
        ws = wb.create_sheet(title)
        ws.append(e["headers"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F2747")
        for row in e["rows"]:
            ws.append([_xl(v) for v in row])
        ws.freeze_panes = "A2"
        for i, h in enumerate(e["headers"], start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(len(h) + 4, 12), 40)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="company_{company_id}.xlsx"'})


def _xl(v):
    from datetime import datetime, time
    from decimal import Decimal
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, datetime) and v.tzinfo:
        return v.replace(tzinfo=None)
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v
