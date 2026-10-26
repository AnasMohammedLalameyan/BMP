"""
check_traceability.py — Verifies that 100% of data sheet rows from
BMP_Master Checklist_Dashboard Requirements.xlsx are mapped in docs/FIELD_TRACEABILITY.md.
Exits with 0 on complete coverage, non-zero if any row is unmapped or missing.
"""

import sys
import openpyxl
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
XLSX_PATH = BASE_DIR / "BMP_Master Checklist_Dashboard Requirements.xlsx"
MD_PATH = BASE_DIR / "docs" / "FIELD_TRACEABILITY.md"


def main():
    if not XLSX_PATH.exists():
        print(f"ERROR: Excel file not found at {XLSX_PATH}")
        sys.exit(1)

    if not MD_PATH.exists():
        print(f"ERROR: Traceability document not found at {MD_PATH}")
        sys.exit(1)

    md_content = MD_PATH.read_text(encoding="utf-8")
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    data_sheets = [
        s for s in wb.sheetnames
        if s not in ['Read Me', "What's Included", 'When Something Becomes Required', 'Standard Lists']
    ]

    missing = []
    total = 0

    for sheet_name in data_sheets:
        ws = wb[sheet_name]
        header_row = None
        for r in range(1, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if vals and (vals[0] == '#' or (len(vals) > 1 and vals[1] == 'Information needed')):
                header_row = r
                continue
            if header_row and r > header_row:
                num_val = vals[0]
                info_val = vals[1] if len(vals) > 1 else None
                if num_val is None and info_val is None:
                    continue
                total += 1
                
                # Check that this sheet and row number appear in the markdown table
                expected_prefix = f"| {sheet_name} | {num_val} |"
                if expected_prefix not in md_content:
                    missing.append((sheet_name, num_val, info_val))

    if missing:
        print(f"FAILED: {len(missing)} out of {total} fields are missing from {MD_PATH.name}:")
        for sheet, row, info in missing:
            print(f"  - {sheet} row #{row}: {info}")
        sys.exit(1)
    else:
        print(f"SUCCESS: All {total} requirement rows across {len(data_sheets)} data sheets are mapped in {MD_PATH.name} (100.0% coverage).")
        sys.exit(0)


if __name__ == "__main__":
    main()
